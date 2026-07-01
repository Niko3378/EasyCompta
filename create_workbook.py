#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Générateur du classeur Excel de gestion comptable.
Produit : EasyCompta.xlsx
"""

import os
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, LineChart, Reference

# ── Palette couleurs ──────────────────────────────────────────────────────────
C_HEADER      = "1F4E79"   # Bleu marine
C_SUB_HEADER  = "2E75B6"   # Bleu moyen
C_ACCENT_L    = "BDD7EE"   # Bleu pâle
C_BLANC       = "FFFFFF"
C_GRIS        = "F2F2F2"
C_GRIS_BD     = "CCCCCC"
C_VERT        = "375623"
C_ORANGE      = "C55A11"
C_VIOLET      = "7030A0"
C_ROUGE       = "C00000"
C_LIGNE_PAIR  = "EBF3FB"

ANNEE = datetime.datetime.now().year


# ── Helpers de style ─────────────────────────────────────────────────────────

def _bd(style='thin', color=C_GRIS_BD):
    return Side(style=style, color=color)

def _border(style='thin', color=C_GRIS_BD):
    s = _bd(style, color)
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type='solid')

def _font(bold=False, size=10, color="000000", italic=False):
    return Font(name='Calibri', bold=bold, size=size, color=color, italic=italic)

def _align(h='left', v='center', wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)

def cell_header(ws, row, col, text, bg=C_HEADER, fg=C_BLANC, size=10, bold=True,
                h_align='center', merge_to=None, row_height=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font    = _font(bold=bold, size=size, color=fg)
    c.fill    = _fill(bg)
    c.alignment = _align(h=h_align, v='center', wrap=True, indent=1 if h_align == 'left' else 0)
    c.border  = _border('thin', C_BLANC)
    if merge_to:
        ws.merge_cells(f"{get_column_letter(col)}{row}:{merge_to}{row}")
    if row_height:
        ws.row_dimensions[row].height = row_height
    return c

def cell_label(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font      = _font(bold=True, size=10)
    c.alignment = _align(h='right', v='center')
    return c

def cell_value_box(ws, row, col, value=None, number_fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font    = _font(size=10)
    c.fill    = _fill(C_GRIS)
    c.border  = _border('thin', "AAAAAA")
    c.alignment = _align(h='left', v='center', indent=1)
    if number_fmt:
        c.number_format = number_fmt
    return c


# ── Feuille base de données (Fournisseurs / Clients) ─────────────────────────

def creer_feuille_db(wb, nom, tab_color, table_name):
    ws = wb.create_sheet(nom)
    ws.sheet_properties.tabColor = tab_color
    ws.freeze_panes = 'A2'

    headers = [
        "ID_Facture", "Fournisseur / Client", "Numéro_Facture",
        "Date_Emission", "Montant_HT", "Montant_TVA", "Montant_TTC",
        "Taux_TVA", "Catégorie", "Lien_PDF",
        "Statut_Paiement", "Date_Import", "Source_Extraction"
    ]
    widths = [16, 28, 20, 14, 14, 14, 14, 11, 20, 45, 16, 14, 18]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell_header(ws, 1, col, h, row_height=38)
        ws.column_dimensions[get_column_letter(col)].width = w

    # Table Excel structurée
    tbl = Table(displayName=table_name,
                ref=f"A1:{get_column_letter(len(headers))}2000")
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True, showFirstColumn=False,
        showLastColumn=False, showColumnStripes=False
    )
    ws.add_table(tbl)

    # Formats numériques lignes 2–2000
    for row in range(2, 2001):
        for col in [5, 6, 7]:                       # HT, TVA, TTC
            ws.cell(row=row, column=col).number_format = '#,##0.00 €'
        ws.cell(row=row, column=8).number_format  = '0.0'   # Taux TVA
        ws.cell(row=row, column=4).number_format  = 'DD/MM/YYYY'
        ws.cell(row=row, column=12).number_format = 'DD/MM/YYYY'

    # Validations
    dv_statut = DataValidation(type="list",
        formula1='"Payé,En attente,En retard,Annulé"',
        allow_blank=True, showErrorMessage=True,
        errorTitle="Statut invalide",
        error="Choisissez : Payé | En attente | En retard | Annulé")
    dv_statut.sqref = "K2:K2000"
    ws.add_data_validation(dv_statut)

    dv_tva = DataValidation(type="list",
        formula1='"0,5.5,10,20"',
        allow_blank=True, showErrorMessage=True,
        errorTitle="Taux invalide",
        error="Taux TVA français : 0 | 5.5 | 10 | 20")
    dv_tva.sqref = "H2:H2000"
    ws.add_data_validation(dv_tva)

    dv_source = DataValidation(type="list",
        formula1='"Manuel,Automatique"', allow_blank=True)
    dv_source.sqref = "M2:M2000"
    ws.add_data_validation(dv_source)

    return ws


# ── Feuille PDF_Import ────────────────────────────────────────────────────────

def creer_feuille_pdf_import(wb):
    ws = wb.create_sheet("PDF_Import")
    ws.sheet_properties.tabColor = C_ORANGE

    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 48
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 22

    # Titre
    cell_header(ws, 1, 1, "IMPORTATION DE FACTURE PDF",
                bg=C_HEADER, size=16, merge_to="D", row_height=48)

    # ─ Section 1 : sélection fichier ─
    cell_header(ws, 3, 1, "1.  SÉLECTIONNER LE FICHIER PDF",
                bg=C_SUB_HEADER, size=12, h_align='left', merge_to="D", row_height=30)

    ws.row_dimensions[4].height = 8  # espace

    cell_label(ws, 5, 1, "Fichier PDF :")
    ws.merge_cells('B5:C5')
    c = ws['B5']
    c.font      = _font(size=10, color="0070C0")
    c.alignment = _align(h='left', v='center', indent=1)
    c.border    = _border('thin', "AAAAAA")
    c.fill      = _fill(C_GRIS)
    ws.row_dimensions[5].height = 28
    # bouton D5 → positionné par VBA

    ws['D5'] = "← Bouton créé par VBA"
    ws['D5'].font = _font(size=9, color="AAAAAA", italic=True)
    ws['D5'].alignment = _align(h='center', v='center')

    ws.row_dimensions[6].height = 8

    # ─ Section 2 : données extraites ─
    cell_header(ws, 7, 1, "2.  DONNÉES EXTRAITES AUTOMATIQUEMENT",
                bg=C_SUB_HEADER, size=12, h_align='left', merge_to="D", row_height=30)

    champs = [
        ("Numéro de facture :",  "B9",  None),
        ("Date d'émission :",    "B10", 'DD/MM/YYYY'),
        ("Montant HT (€) :",     "B11", '#,##0.00'),
        ("Montant TVA (€) :",    "B12", '#,##0.00'),
        ("Montant TTC (€) :",    "B13", '#,##0.00'),
        ("Taux TVA (%) :",       "B14", '0.0'),
        ("Type de facture :",    "B15", None),
    ]
    for i, (label, ref, fmt) in enumerate(champs, 9):
        cell_label(ws, i, 1, label)
        ws.merge_cells(f'B{i}:C{i}')
        cell_value_box(ws, i, 2, number_fmt=fmt)
        ws.row_dimensions[i].height = 26

    ws.row_dimensions[16].height = 8

    # ─ Section 3 : actions ─
    cell_header(ws, 17, 1, "3.  ACTIONS",
                bg=C_SUB_HEADER, size=12, h_align='left', merge_to="D", row_height=30)
    ws['A18'] = "→ Les boutons ci-dessous sont créés automatiquement par la macro VBA Setup_Boutons"
    ws['A18'].font = _font(size=9, color="888888", italic=True)
    ws.merge_cells('A18:D18')
    ws.row_dimensions[18].height = 22
    ws.row_dimensions[19].height = 40  # zone boutons
    ws.row_dimensions[20].height = 40

    ws.row_dimensions[21].height = 8

    # ─ Zone statut ─
    cell_header(ws, 22, 1, "STATUT",
                bg=C_SUB_HEADER, size=11, h_align='left', merge_to="D", row_height=26)
    ws.merge_cells('A23:D25')
    ws['A23'] = "En attente d'importation..."
    ws['A23'].font = _font(size=10, color="666666", italic=True)
    ws['A23'].alignment = _align(h='left', v='top', wrap=True, indent=1)
    ws['A23'].fill = _fill(C_GRIS)
    ws['A23'].border = _border('thin', "BBBBBB")
    for r in [23, 24, 25]:
        ws.row_dimensions[r].height = 22

    return ws


# ── Feuille Pivot_Analytics ───────────────────────────────────────────────────

def creer_feuille_pivot(wb):
    ws = wb.create_sheet("Pivot_Analytics")
    ws.sheet_properties.tabColor = "375623"
    ws.freeze_panes = 'B5'

    # Largeurs colonnes A–F
    for col, w in enumerate([16, 22, 24, 20, 20, 18], 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    # Colonnes graphiques
    for col in range(7, 22):
        ws.column_dimensions[get_column_letter(col)].width = 11

    # Titre
    cell_header(ws, 1, 1, "TABLEAU DE BORD ANALYTIQUE",
                bg=C_HEADER, size=16, merge_to="F", row_height=48)

    # Sélecteur d'année (référence commune)
    ws['A3'] = "Année affichée :"
    ws['A3'].font = _font(bold=True, size=11)
    ws['A3'].alignment = _align(h='right', v='center')
    ws['B3'] = f'=TVA_Mensuelle!B3'
    ws['B3'].font = _font(bold=True, size=14, color=C_HEADER)
    ws['B3'].alignment = _align(h='center', v='center')
    ws['B3'].border = _border('double', C_HEADER)
    ws['C3'] = "(modifiable dans TVA_Mensuelle, cellule B3)"
    ws['C3'].font = _font(size=9, color="888888", italic=True)
    ws.merge_cells('C3:F3')
    ws.row_dimensions[3].height = 30

    # ── Tableau synthèse mensuelle ─────────────────────────────────────
    hdrs = ["Mois", "CA HT\n(Clients)", "Dépenses HT\n(Fournisseurs)",
            "TVA Collectée", "TVA Déductible", "Solde TVA"]
    for col, h in enumerate(hdrs, 1):
        cell_header(ws, 4, col, h, bg="366092", row_height=38)

    MOIS = ["Janvier","Février","Mars","Avril","Mai","Juin",
            "Juillet","Août","Septembre","Octobre","Novembre","Décembre"]

    for i, mois in enumerate(MOIS, 1):
        r = 4 + i
        # Nom mois
        c = ws.cell(row=r, column=1, value=mois)
        c.font = _font(bold=True, size=10)
        c.alignment = _align(indent=1)
        c.border = _border()

        # Formules SUMPRODUCT avec IFERROR pour gérer cellules vides
        ann = "TVA_Mensuelle!$B$3"
        m = i

        def sp_client(col_name):
            return (f'=IFERROR(SUMPRODUCT('
                    f'(MONTH(Clients_DB[Date_Emission])={m})*'
                    f'(YEAR(Clients_DB[Date_Emission])={ann})*'
                    f'Clients_DB[{col_name}]),0)')

        def sp_fourn(col_name):
            return (f'=IFERROR(SUMPRODUCT('
                    f'(MONTH(Fournisseurs_DB[Date_Emission])={m})*'
                    f'(YEAR(Fournisseurs_DB[Date_Emission])={ann})*'
                    f'Fournisseurs_DB[{col_name}]),0)')

        formules = [
            sp_client("Montant_HT"),
            sp_fourn("Montant_HT"),
            sp_client("Montant_TVA"),
            sp_fourn("Montant_TVA"),
        ]
        for col, f in enumerate(formules, 2):
            c = ws.cell(row=r, column=col, value=f)
            c.number_format = '#,##0.00 €'
            c.border = _border()
            c.alignment = _align(h='right', v='center')

        # Solde TVA = collectée - déductible
        c = ws.cell(row=r, column=6, value=f'=D{r}-E{r}')
        c.number_format = '#,##0.00 €'
        c.border = _border()
        c.alignment = _align(h='right', v='center')

        # Alternance couleur
        if i % 2 == 0:
            for col in range(1, 7):
                ws.cell(row=r, column=col).fill = _fill(C_LIGNE_PAIR)
        ws.row_dimensions[r].height = 22

    # Ligne total
    tr = 17
    c0 = ws.cell(row=tr, column=1, value="TOTAL ANNUEL")
    c0.font = _font(bold=True, size=10, color=C_BLANC)
    c0.fill = _fill(C_HEADER)
    c0.alignment = _align(h='center')
    c0.border = _border('thin', C_BLANC)
    for col in range(2, 7):
        cl = get_column_letter(col)
        c = ws.cell(row=tr, column=col, value=f'=SUM({cl}5:{cl}16)')
        c.number_format = '#,##0.00 €'
        c.font = _font(bold=True, size=10, color=C_BLANC)
        c.fill = _fill(C_HEADER)
        c.border = _border('thin', C_BLANC)
        c.alignment = _align(h='right')
    ws.row_dimensions[tr].height = 28

    # ── Graphique 1 : CA vs Dépenses ──────────────────────────────────
    chart1 = BarChart()
    chart1.type    = "col"
    chart1.title   = f"CA vs Dépenses mensuels HT — {ANNEE}"
    chart1.y_axis.title = "Montant (€)"
    chart1.x_axis.title = "Mois"
    chart1.style   = 10
    chart1.width   = 22
    chart1.height  = 13

    data1 = Reference(ws, min_col=2, max_col=3, min_row=4, max_row=16)
    cats1 = Reference(ws, min_col=1, min_row=5, max_row=16)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    ws.add_chart(chart1, "H3")

    # ── Graphique 2 : TVA collectée vs déductible ─────────────────────
    chart2 = BarChart()
    chart2.type    = "col"
    chart2.title   = f"TVA collectée vs TVA déductible — {ANNEE}"
    chart2.y_axis.title = "Montant (€)"
    chart2.x_axis.title = "Mois"
    chart2.style   = 11
    chart2.width   = 22
    chart2.height  = 13

    data2 = Reference(ws, min_col=4, max_col=5, min_row=4, max_row=16)
    cats2 = Reference(ws, min_col=1, min_row=5, max_row=16)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    ws.add_chart(chart2, "H22")

    # ── Graphique 3 : Solde TVA mensuel (courbe) ──────────────────────
    chart3 = LineChart()
    chart3.title   = f"Évolution solde TVA — {ANNEE}"
    chart3.y_axis.title = "Solde (€)"
    chart3.x_axis.title = "Mois"
    chart3.style   = 12
    chart3.width   = 22
    chart3.height  = 13

    data3 = Reference(ws, min_col=6, max_col=6, min_row=4, max_row=16)
    cats3 = Reference(ws, min_col=1, min_row=5, max_row=16)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    ws.add_chart(chart3, "H41")

    return ws


# ── Feuille TVA_Mensuelle ─────────────────────────────────────────────────────

def creer_feuille_tva(wb):
    ws = wb.create_sheet("TVA_Mensuelle")
    ws.sheet_properties.tabColor = C_ROUGE
    ws.freeze_panes = 'A6'

    for col, w in enumerate([20, 12, 15, 22, 22, 20, 20, 25], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # Titre
    cell_header(ws, 1, 1, "CALCUL DE TVA MENSUELLE",
                bg=C_HEADER, size=16, merge_to="H", row_height=48)

    # Sélecteur année
    ws['A3'] = "Année fiscale :"
    ws['A3'].font = _font(bold=True, size=12)
    ws['A3'].alignment = _align(h='right', v='center')

    ws['B3'] = ANNEE
    ws['B3'].font = _font(bold=True, size=16, color=C_HEADER)
    ws['B3'].alignment = _align(h='center', v='center')
    ws['B3'].border = _border('double', C_HEADER)

    ws['C3'] = "← Saisissez l'année souhaitée"
    ws['C3'].font = _font(size=9, italic=True, color="888888")
    ws.merge_cells('C3:H3')
    ws.row_dimensions[3].height = 32

    ws.row_dimensions[4].height = 8

    # En-têtes tableau TVA
    hdrs = ["Mois", "Année", "TVA Collectée\n(Clients)",
            "TVA Déductible\n(Fournisseurs)", "TVA à Payer\n(Collectée − Déductible)",
            "Solde\nCumulé", "Statut", "Échéance\nDéclaration"]
    for col, h in enumerate(hdrs, 1):
        cell_header(ws, 5, col, h, bg=C_HEADER, row_height=42)

    MOIS = ["Janvier","Février","Mars","Avril","Mai","Juin",
            "Juillet","Août","Septembre","Octobre","Novembre","Décembre"]

    for i, mois in enumerate(MOIS, 1):
        r = 5 + i
        m = i

        ws.cell(row=r, column=1, value=mois).font = _font(bold=True, size=10)
        ws.cell(row=r, column=1).alignment = _align(indent=1)

        ws.cell(row=r, column=2, value='=$B$3')
        ws.cell(row=r, column=2).number_format = '0'
        ws.cell(row=r, column=2).alignment = _align(h='center')

        # TVA collectée clients
        f_tc = (f'=IFERROR(SUMPRODUCT('
                f'(MONTH(Clients_DB[Date_Emission])={m})*'
                f'(YEAR(Clients_DB[Date_Emission])=$B$3)*'
                f'Clients_DB[Montant_TVA]),0)')
        ws.cell(row=r, column=3, value=f_tc).number_format = '#,##0.00 €'

        # TVA déductible fournisseurs
        f_td = (f'=IFERROR(SUMPRODUCT('
                f'(MONTH(Fournisseurs_DB[Date_Emission])={m})*'
                f'(YEAR(Fournisseurs_DB[Date_Emission])=$B$3)*'
                f'Fournisseurs_DB[Montant_TVA]),0)')
        ws.cell(row=r, column=4, value=f_td).number_format = '#,##0.00 €'

        # TVA à payer
        ws.cell(row=r, column=5, value=f'=C{r}-D{r}').number_format = '#,##0.00 €'

        # Solde cumulé
        if i == 1:
            ws.cell(row=r, column=6, value=f'=E{r}').number_format = '#,##0.00 €'
        else:
            ws.cell(row=r, column=6, value=f'=F{r-1}+E{r}').number_format = '#,##0.00 €'

        # Statut textuel
        ws.cell(row=r, column=7,
                value=f'=IF(E{r}>0,"À PAYER : "&TEXT(E{r},"#,##0.00 €"),'
                      f'IF(E{r}<0,"CRÉDIT : "&TEXT(ABS(E{r}),"#,##0.00 €"),"Équilibré"))')
        ws.cell(row=r, column=7).alignment = _align(h='center')

        # Échéance (le 19 du mois suivant)
        if m == 12:
            date_ech = f'=DATE($B$3+1,1,19)'
        else:
            date_ech = f'=DATE($B$3,{m+1},19)'
        ws.cell(row=r, column=8, value=date_ech).number_format = 'DD/MM/YYYY'
        ws.cell(row=r, column=8).alignment = _align(h='center')

        # Style
        for col in range(1, 9):
            ws.cell(row=r, column=col).border = _border()
            ws.cell(row=r, column=col).alignment = _align(
                h='center' if col > 1 else 'left',
                v='center',
                indent=1 if col == 1 else 0
            )
            if i % 2 == 0:
                ws.cell(row=r, column=col).fill = _fill(C_LIGNE_PAIR)
        ws.row_dimensions[r].height = 24

    # Ligne total annuel
    tr = 18
    c = ws.cell(row=tr, column=1, value="TOTAL ANNUEL")
    c.font = _font(bold=True, size=11, color=C_BLANC)
    c.fill = _fill(C_HEADER)
    c.alignment = _align(h='center')
    c.border = _border('thin', C_BLANC)

    for col in range(3, 7):
        cl = get_column_letter(col)
        cell = ws.cell(row=tr, column=col,
                       value=f'=SUM({cl}6:{cl}17)' if col < 6 else f'=F17')
        cell.number_format = '#,##0.00 €'
        cell.font = _font(bold=True, size=11, color=C_BLANC)
        cell.fill = _fill(C_HEADER)
        cell.border = _border('thin', C_BLANC)
        cell.alignment = _align(h='center')
    ws.row_dimensions[tr].height = 30

    # ── Synthèse par taux TVA ─────────────────────────────────────────
    ws.row_dimensions[20].height = 8
    cell_header(ws, 21, 1, "SYNTHÈSE PAR TAUX DE TVA",
                bg=C_SUB_HEADER, size=12, h_align='left', merge_to="H", row_height=30)

    hdrs_taux = ["Taux TVA", "TVA Collectée (Clients)",
                 "TVA Déductible (Fournisseurs)", "Solde Net"]
    for col, h in enumerate(hdrs_taux, 1):
        cell_header(ws, 22, col, h, bg="366092", row_height=28)

    taux_list = [("0 %", 0), ("5,5 %", 5.5), ("10 %", 10), ("20 %", 20)]
    for j, (label, taux) in enumerate(taux_list, 1):
        r = 22 + j
        ws.cell(row=r, column=1, value=label).font = _font(bold=True, size=10)
        ws.cell(row=r, column=1).alignment = _align(h='center')

        f_c = (f'=IFERROR(SUMPRODUCT('
               f'(Clients_DB[Taux_TVA]={taux})*'
               f'(YEAR(Clients_DB[Date_Emission])=$B$3)*'
               f'Clients_DB[Montant_TVA]),0)')
        ws.cell(row=r, column=2, value=f_c).number_format = '#,##0.00 €'

        f_d = (f'=IFERROR(SUMPRODUCT('
               f'(Fournisseurs_DB[Taux_TVA]={taux})*'
               f'(YEAR(Fournisseurs_DB[Date_Emission])=$B$3)*'
               f'Fournisseurs_DB[Montant_TVA]),0)')
        ws.cell(row=r, column=3, value=f_d).number_format = '#,##0.00 €'

        ws.cell(row=r, column=4, value=f'=B{r}-C{r}').number_format = '#,##0.00 €'

        for col in range(1, 5):
            ws.cell(row=r, column=col).border = _border()
            ws.cell(row=r, column=col).alignment = _align(
                h='center' if col != 1 else 'center')
            if j % 2 == 0:
                ws.cell(row=r, column=col).fill = _fill(C_LIGNE_PAIR)
        ws.row_dimensions[r].height = 24

    return ws


# ── Feuille Guide_Utilisateur ─────────────────────────────────────────────────

def creer_feuille_guide(wb):
    ws = wb.create_sheet("Guide_Utilisateur")
    ws.sheet_properties.tabColor = C_VIOLET

    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 80
    ws.column_dimensions['C'].width = 30

    cell_header(ws, 1, 1, "GUIDE UTILISATEUR — SYSTÈME DE GESTION COMPTABLE IMC",
                bg=C_VIOLET, size=16, merge_to="C", row_height=52)

    contenu = [
        # (row, texte, is_section, bg_couleur)
        (3,  "INSTALLATION INITIALE",                                   True,  C_SUB_HEADER),
        (4,  "1.  Exécutez install.bat en tant qu'administrateur.",      False, None),
        (5,  "     Ce script installe Python, les bibliothèques et crée le fichier Excel.", False, None),
        (6,  "2.  Ouvrez EasyCompta.xlsx dans Excel 365.",        False, None),
        (7,  "3.  Enregistrez-le en .xlsm : Fichier → Enregistrer sous → "
             "Classeur Excel prenant en charge les macros (*.xlsm).",    False, None),
        (8,  "4.  Ouvrez l'éditeur VBA : Alt + F11.",                   False, None),
        (9,  "5.  Importez le module : Fichier → Importer → vba_Comptabilite.bas.", False, None),
        (10, "6.  Fermez l'éditeur, puis lancez : Alt + F8 → Setup_Boutons → Exécuter.", False, None),
        (12, "IMPORTER UNE FACTURE PDF (AUTOMATIQUE)",                   True,  C_SUB_HEADER),
        (13, "1.  Allez sur la feuille PDF_Import.",                     False, None),
        (14, "2.  Cliquez sur « Importer PDF » et sélectionnez votre fichier.", False, None),
        (15, "3.  Cliquez sur « Extraire données » — Python analyse le PDF.", False, None),
        (16, "4.  Vérifiez les champs extraits (corrigez si besoin).",   False, None),
        (17, "5.  Cliquez sur « Envoyer vers base » pour enregistrer.",  False, None),
        (18, "     Le système détecte automatiquement Fournisseur ou Client.", False, None),
        (20, "SAISIE MANUELLE",                                          True,  C_SUB_HEADER),
        (21, "1.  Ouvrez Fournisseurs_DB ou Clients_DB.",                False, None),
        (22, "2.  Saisissez directement dans les colonnes.",             False, None),
        (23, "3.  Utilisez les listes déroulantes pour Statut_Paiement et Taux_TVA.", False, None),
        (24, "4.  Laissez ID_Facture vide : la macro le génère automatiquement.", False, None),
        (26, "CALCUL TVA",                                               True,  C_SUB_HEADER),
        (27, "1.  Feuille TVA_Mensuelle : modifiez la cellule B3 pour changer d'année.", False, None),
        (28, "2.  Les calculs se mettent à jour automatiquement.",       False, None),
        (29, "3.  TVA à payer = TVA collectée (Clients) − TVA déductible (Fournisseurs).", False, None),
        (30, "4.  L'échéance légale est le 19 du mois suivant (régime mensuel).", False, None),
        (32, "TABLEAUX DE BORD (Pivot_Analytics)",                       True,  C_SUB_HEADER),
        (33, "•  Graphiques automatiquement mis à jour à chaque modification.", False, None),
        (34, "•  L'année affichée suit la cellule B3 de TVA_Mensuelle.", False, None),
        (35, "•  Utilisez le bouton « Rafraîchir » après un import groupé.", False, None),
        (37, "TAUX DE TVA FRANÇAIS",                                     True,  "C55A11"),
        (38, "•  0 %    — Exonération (exportations, certains services de santé).", False, None),
        (39, "•  5,5 % — Taux réduit (produits alimentaires, livres, abonnements énergie).", False, None),
        (40, "•  10 %  — Taux intermédiaire (restauration, transport, travaux).", False, None),
        (41, "•  20 %  — Taux normal (majorité des biens et services).",  False, None),
        (43, "LIMITES DE L'EXTRACTION PDF AUTOMATIQUE",                  True,  C_ROUGE),
        (44, "•  PDFs scannés (images) : non supportés sans module OCR.", False, None),
        (45, "•  Formats très non-standard : extraction partielle possible.", False, None),
        (46, "•  Toujours vérifier les données avant envoi en base.",    False, None),
        (48, "FICHIERS DU PROJET",                                       True,  C_VIOLET),
        (49, "•  create_workbook.py    — génère la structure Excel.",    False, None),
        (50, "•  pdf_extractor.py      — moteur d'extraction PDF (Python).", False, None),
        (51, "•  vba_Comptabilite.bas  — macros VBA à importer dans Excel.", False, None),
        (52, "•  install.bat           — script d'installation Windows.", False, None),
        (53, "•  requirements.txt      — dépendances Python.",           False, None),
    ]

    for row, texte, is_section, bg in contenu:
        c = ws.cell(row=row, column=2, value=texte)
        if is_section:
            c.font = _font(bold=True, size=12, color=C_BLANC)
            c.fill = _fill(bg)
            c.alignment = _align(h='left', v='center', indent=1)
            ws.row_dimensions[row].height = 28
            # Étendre sur colonne C aussi
            ws.merge_cells(f'B{row}:C{row}')
        else:
            c.font = _font(size=10)
            c.alignment = _align(h='left', v='center', indent=2)
            ws.row_dimensions[row].height = 19

    return ws


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    sortie = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          "EasyCompta.xlsx")

    print("Creation du classeur Excel...")
    wb = Workbook()
    wb.remove(wb.active)   # supprimer la feuille vide par défaut

    print("  -> Fournisseurs_DB")
    creer_feuille_db(wb, "Fournisseurs_DB", C_HEADER,  "Fournisseurs_DB")

    print("  -> Clients_DB")
    creer_feuille_db(wb, "Clients_DB",      C_VERT,    "Clients_DB")

    print("  -> PDF_Import")
    creer_feuille_pdf_import(wb)

    print("  -> Pivot_Analytics")
    creer_feuille_pivot(wb)

    print("  -> TVA_Mensuelle")
    creer_feuille_tva(wb)

    print("  -> Guide_Utilisateur")
    creer_feuille_guide(wb)

    wb.save(sortie)
    print("\nFichier cree : " + sortie)
    print("\nProchaines etapes :")
    print("  1. Ouvrez le fichier dans Excel")
    print("  2. Enregistrez-le en .xlsm (Fichier -> Enregistrer sous)")
    print("  3. Alt+F11 -> Fichier -> Importer -> vba_Comptabilite.bas")
    print("  4. Alt+F8 -> Setup_Boutons -> Executer")


if __name__ == '__main__':
    main()
